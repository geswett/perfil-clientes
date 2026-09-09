import os
import tempfile
import traceback

import docx as docx_reader
from flask import Flask, request, jsonify, render_template, send_file

import claude_service
from docx_generator import generar_docx

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB (fotos y archivos de transcripción)

ALLOWED_TRANSCRIPCION_EXT = {".doc", ".docx", ".pdf", ".txt"}
ALLOWED_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif"}

MIME_BY_EXT = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".gif": "image/gif",
}


def extraer_texto_de_archivo(tmp_path: str, ext: str) -> str:
    """Extrae el texto plano de un archivo de transcripción ya hecho (.doc, .docx, .pdf o .txt)."""
    if ext == ".txt":
        with open(tmp_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read().strip()

    elif ext == ".docx":
        documento = docx_reader.Document(tmp_path)
        return "\n".join(p.text for p in documento.paragraphs if p.text.strip())

    elif ext == ".doc":
        # .doc es el formato binario antiguo de Word (previo a 2007) — no hay forma confiable
        # de leerlo sin instalar LibreOffice en el servidor. Lo intentamos igual por si el
        # archivo en realidad es un .docx mal nombrado; si no, pedimos que lo conviertan.
        try:
            documento = docx_reader.Document(tmp_path)
            return "\n".join(p.text for p in documento.paragraphs if p.text.strip())
        except Exception as e:
            raise ValueError(
                "No se pudo leer este archivo .doc (formato antiguo de Word, previo a 2007). "
                "Ábrelo en Word o Google Docs y guárdalo/expórtalo como .docx o PDF, y vuelve a "
                "subirlo."
            ) from e

    elif ext == ".pdf":
        return claude_service.extract_pdf_text(tmp_path)

    raise ValueError(f"Extensión no soportada para transcripción: {ext}")


def _perfil_parece_vacio(perfil: dict) -> bool:
    """Chequeo de seguridad: si Claude devuelve el Perfil de Cargo con las secciones
    principales vacías (por ejemplo, por un problema con el schema o con la respuesta de la
    API), es mejor avisar con un error claro que mostrar un formulario en blanco sin
    explicación."""
    secciones_objeto = ["empresa", "organigrama", "descripcion_cargo", "condiciones_laborales"]
    if any(not perfil.get(s) for s in secciones_objeto):
        return True
    if not perfil.get("requisitos") or not perfil.get("competencias"):
        return True
    if not (perfil.get("empresa") or {}).get("definicion_empresa"):
        return True
    return False


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/healthz")
def healthz():
    return {"status": "ok"}


@app.route("/api/procesar", methods=["POST"])
def procesar():
    """Recibe un archivo de transcripción, una foto de notas, o texto escrito, y devuelve el
    Perfil de Cargo estructurado con Claude."""
    modo = request.form.get("modo")
    empresa = request.form.get("empresa", "").strip()
    cargo = request.form.get("cargo", "").strip()

    if modo not in ("transcripcion", "imagen", "texto"):
        return jsonify({"error": "Modo inválido. Debe ser 'transcripcion', 'imagen' o 'texto'."}), 400

    tmp_path = None
    try:
        if modo == "texto":
            raw_text = request.form.get("texto", "").strip()
            if not raw_text:
                return jsonify({"error": "No se recibió texto para procesar."}), 400

        elif modo == "transcripcion":
            archivo = request.files.get("archivo")
            if not archivo or not archivo.filename:
                return jsonify({"error": "Sube el archivo de la transcripción (.doc, .docx, .pdf o .txt)."}), 400

            ext = os.path.splitext(archivo.filename)[1].lower()
            if ext not in ALLOWED_TRANSCRIPCION_EXT:
                return jsonify({
                    "error": f"Formato '{ext}' no soportado para transcripción. "
                             f"Formatos permitidos: {', '.join(sorted(ALLOWED_TRANSCRIPCION_EXT))}"
                }), 400

            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                archivo.save(tmp.name)
                tmp_path = tmp.name

            try:
                raw_text = extraer_texto_de_archivo(tmp_path, ext)
            except ValueError as e:
                return jsonify({"error": str(e)}), 400

            if not raw_text.strip():
                return jsonify({"error": "La transcripción está vacía."}), 400

        else:  # modo == "imagen"
            archivo = request.files.get("archivo")
            if not archivo or archivo.filename == "":
                return jsonify({"error": "No se recibió ningún archivo."}), 400

            ext = os.path.splitext(archivo.filename)[1].lower()
            if ext not in ALLOWED_IMAGE_EXT:
                return jsonify({
                    "error": f"Formato '{ext}' no soportado para fotos. "
                             f"Formatos permitidos: {', '.join(sorted(ALLOWED_IMAGE_EXT))}"
                }), 400

            mime_type = MIME_BY_EXT.get(ext, "image/jpeg")

            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                archivo.save(tmp.name)
                tmp_path = tmp.name

            raw_text = claude_service.transcribe_image(tmp_path, mime_type)

        perfil = claude_service.structure_profile(raw_text, empresa=empresa, cargo=cargo)

        if _perfil_parece_vacio(perfil):
            return jsonify({
                "error": "Claude leyó el contenido pero no logró estructurar el Perfil de Cargo "
                         "correctamente (quedó vacío). Prueba de nuevo — normalmente basta con "
                         "reintentar. Si vuelve a pasar, avísale a quien mantiene esta herramienta."
            }), 502

        return jsonify({
            "transcripcion": raw_text,
            "perfil": perfil,
        })

    except RuntimeError as e:
        return jsonify({"error": str(e)}), 500
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": f"Ocurrió un error procesando la solicitud: {e}"}), 500
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)


@app.route("/api/generar-docx", methods=["POST"])
def api_generar_docx():
    """Recibe el perfil (ya revisado/editado por el usuario) y devuelve el .docx."""
    data = request.get_json(force=True, silent=True) or {}
    perfil = data.get("perfil")
    empresa = data.get("empresa", "")
    cargo = data.get("cargo", "")
    consultor = data.get("consultor", "")

    if not perfil:
        return jsonify({"error": "Falta el perfil a exportar."}), 400

    try:
        contenido, filename = generar_docx(perfil, empresa, cargo, consultor)
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": f"No se pudo generar el documento: {e}"}), 500

    import io
    return send_file(
        io.BytesIO(contenido),
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
